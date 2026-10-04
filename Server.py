import base64
import json
import os
import re
import time
from datetime import datetime
from io import BytesIO

from dotenv import load_dotenv
from flask import Flask, jsonify, request, send_from_directory
from openai import OpenAI

# Load secrets from .env in this folder. The key never goes to the browser.
load_dotenv()

OPENAI_MODEL = os.getenv("OPENAI_MODEL", "gpt-4o-mini")
MAX_FILE_BYTES = 12 * 1024 * 1024
MAX_FILES = 12
MAX_TEXT_CHARS = 120000
MAX_REQUEST_BYTES = 25 * 1024 * 1024   # whole upload, so a huge request cannot fill memory
RATE_MAX = int(os.getenv("RATE_MAX", "20"))   # requests per IP per 10 minutes
RATE_WINDOW_SECONDS = 600
PLACEHOLDER_KEYS = {
    "",
    "paste_your_openai_api_key_here",
    "syllabuddy",
    "your_api_key_here",
}

ALLOWED_TYPES = {
    "Class",
    "Assignment",
    "Reading",
    "Quiz",
    "Exam",
    "Project",
}

DATE_PATTERN = re.compile(r"^\d{4}-\d{2}-\d{2}$")
TIME_24 = re.compile(r"^([01]?\d|2[0-3]):([0-5]\d)$")
TIME_AMPM = re.compile(
    r"^\s*(\d{1,2})(?::(\d{2}))?\s*([ap])\.?m\.?\s*$",
    re.I,
)

SYSTEM_INSTRUCTIONS = """
You build one student schedule from whatever they pasted or uploaded, including messy notes like
"algebra 100 at 8 am to 5 pm".

Include:
- class meetings (type Class), such as "Algebra 100 at 8 am to 5 pm"
- assignments, homework, labs, papers, projects
- quizzes, tests, midterms, finals
- readings with dates

Rules:
- type must be one of: Class, Assignment, Reading, Quiz, Exam, Project.
- A line that is only a course name plus a time range (8am to 5pm, 8:00-17:00) is a Class meeting, not an assignment.
- For Class rows, time is the start time as HH:MM. Put the end time in details (example: "Until 17:00").
- Midterm, final, and "test" are Exam. "Quiz" is Quiz. Do not label an exam or quiz as Assignment.
- item is a short title. Do not invent titles.
- due must be YYYY-MM-DD when the text gives a real calendar date. Class meetings often have no date; then due is "".
- If month and day are given but year is not, use the year from Today's date. If that date is invalid (for example February 30), set due to "".
- For assignments, quizzes, exams, and projects, time must be a due time, not a lecture meeting time.
- If the student provided a default due time, you may copy it onto dated Assignment/Quiz/Exam/Project rows with no time. Do not apply it to Class or Reading.
- If date or time cannot be known, leave that field "".
- source must be a short quote copied from the pasted text.
- details is extra scope from the text only (end time, word limit, pages). Otherwise "".
- link is a URL copied from the text only. Otherwise "".
- flag is a short warning if something is unclear. Otherwise "".
- Do not invent items, dates, times, links, or quotes.
- If nothing can be extracted, return {"items": []}.
""".strip()

ITEM_SCHEMA = {
    "type": "object",
    "additionalProperties": False,
    "required": ["items"],
    "properties": {
        "items": {
            "type": "array",
            "items": {
                "type": "object",
                "additionalProperties": False,
                "required": ["type", "item", "due", "time", "source", "details", "link", "flag"],
                "properties": {
                    "type": {
                        "type": "string",
                        "enum": [
                            "Class",
                            "Assignment",
                            "Reading",
                            "Quiz",
                            "Exam",
                            "Project",
                        ],
                    },
                    "item": {"type": "string"},
                    "due": {"type": "string"},
                    "time": {"type": "string"},
                    "source": {"type": "string"},
                    "details": {"type": "string"},
                    "link": {"type": "string"},
                    "flag": {"type": "string"},
                },
            },
        }
    },
}

app = Flask(__name__)
app.config["MAX_CONTENT_LENGTH"] = MAX_REQUEST_BYTES
_hits = {}
ROOT = os.path.dirname(os.path.abspath(__file__))
IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".gif", ".webp", ".bmp"}
TEXT_EXTS = {
    ".txt",
    ".md",
    ".csv",
    ".tsv",
    ".json",
    ".html",
    ".htm",
    ".rtf",
    ".log",
}


def json_error(message, status_code):
    return jsonify({"error": message}), status_code


def rate_limited():
    """Approximate per-IP limit. It resets when the server restarts; the provider spending cap is the real safeguard."""
    forwarded = request.headers.get("X-Forwarded-For", "")
    ip = (forwarded.split(",")[0].strip() if forwarded else request.remote_addr) or "unknown"
    now = time.time()
    recent = [t for t in _hits.get(ip, []) if now - t < RATE_WINDOW_SECONDS]
    recent.append(now)
    _hits[ip] = recent
    return len(recent) > RATE_MAX


def valid_due_date(value):
    if not value:
        return ""
    text = str(value).strip()
    if not DATE_PATTERN.fullmatch(text):
        return ""
    try:
        datetime.strptime(text, "%Y-%m-%d")
    except ValueError:
        return ""
    return text


def valid_time(value):
    if not value:
        return ""
    text = str(value).strip()
    match = TIME_24.fullmatch(text)
    if match:
        return f"{int(match.group(1)):02d}:{match.group(2)}"
    match = TIME_AMPM.fullmatch(text)
    if not match:
        return ""
    hour = int(match.group(1))
    minute = int(match.group(2) or "0")
    period = match.group(3).lower()
    if hour < 1 or hour > 12 or minute > 59:
        return ""
    if period == "p" and hour != 12:
        hour += 12
    if period == "a" and hour == 12:
        hour = 0
    return f"{hour:02d}:{minute:02d}"


def weekday_mismatch(due, text):
    if not due:
        return ""
    match = re.search(
        r"\b(monday|mon|tuesday|tues|tue|wednesday|wed|thursday|thurs|thur|thu|friday|fri|saturday|sat|sunday|sun)\b",
        text,
        re.I,
    )
    if not match:
        return ""
    names = {
        "mon": 0,
        "tue": 1,
        "wed": 2,
        "thu": 3,
        "fri": 4,
        "sat": 5,
        "sun": 6,
    }
    stated = names[match.group(1)[:3].lower()]
    actual = datetime.strptime(due, "%Y-%m-%d").weekday()
    if stated != actual:
        return "weekday does not match date"
    return ""


def appears_in(snippet, syllabus):
    hay = re.sub(r"\s+", " ", (syllabus or "").lower())
    needle = re.sub(r"\s+", " ", (snippet or "").strip().lower())
    return bool(needle) and needle in hay


def clean_items(raw_items, syllabus=""):
    cleaned = []
    for raw in raw_items or []:
        if not isinstance(raw, dict):
            continue
        name = str(raw.get("item") or "").strip()
        if not name:
            continue
        item_type = str(raw.get("type") or "Assignment").strip()
        if item_type not in ALLOWED_TYPES:
            item_type = "Assignment"
        due = valid_due_date(raw.get("due"))
        source = str(raw.get("source") or "").strip()
        details = str(raw.get("details") or "").strip()
        link = str(raw.get("link") or "").strip()
        flags = []
        extra = str(raw.get("flag") or "").strip()
        if extra:
            flags.append(extra)
        if source and syllabus and not appears_in(source, syllabus):
            source = source[:180]
            flags.append("source not copied from paste")
        if details and syllabus and item_type != "Class" and not appears_in(details, syllabus):
            details = ""
        if link:
            if not re.match(r"^https?://", link, re.I) or (
                syllabus and link not in syllabus
            ):
                link = ""
        mismatch = weekday_mismatch(due, name + " " + source)
        if mismatch:
            flags.append(mismatch)
        if not due and item_type != "Class":
            flags.append("no date")
        cleaned.append(
            {
                "type": item_type,
                "item": name,
                "due": due,
                "time": valid_time(raw.get("time")),
                "source": source,
                "details": details,
                "link": link,
                "flag": "; ".join(dict.fromkeys(flags)),
            }
        )
    return cleaned


def decode_bytes(data):
    for encoding in ("utf-8", "utf-16", "latin-1"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def extract_pdf(data):
    from pypdf import PdfReader

    reader = PdfReader(BytesIO(data))
    pages = []
    for page in reader.pages:
        pages.append(page.extract_text() or "")
    return "\n".join(pages).strip()


def extract_docx(data):
    from docx import Document

    document = Document(BytesIO(data))
    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(part for part in parts if part).strip()


def extract_xlsx(data):
    from openpyxl import load_workbook

    workbook = load_workbook(BytesIO(data), data_only=True, read_only=True)
    lines = []
    for sheet in workbook.worksheets:
        lines.append("Sheet: " + sheet.title)
        for row in sheet.iter_rows(values_only=True):
            values = ["" if cell is None else str(cell) for cell in row]
            if any(value.strip() for value in values):
                lines.append("\t".join(values))
    return "\n".join(lines).strip()


def extract_pptx(data):
    from pptx import Presentation

    presentation = Presentation(BytesIO(data))
    parts = []
    for slide in presentation.slides:
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False):
                parts.append(shape.text)
    return "\n".join(part for part in parts if part).strip()


def mime_for_image(ext):
    return {
        ".png": "image/png",
        ".jpg": "image/jpeg",
        ".jpeg": "image/jpeg",
        ".gif": "image/gif",
        ".webp": "image/webp",
        ".bmp": "image/bmp",
    }.get(ext, "image/png")


def read_uploaded_files(files):
    texts = []
    images = []
    names = []
    if len(files) > MAX_FILES:
        raise RuntimeError("Upload at most " + str(MAX_FILES) + " files at a time.")
    for stored in files:
        filename = (stored.filename or "").strip()
        if not filename:
            continue
        data = stored.read()
        if not data:
            continue
        if len(data) > MAX_FILE_BYTES:
            raise RuntimeError(filename + " is larger than 12 MB.")
        names.append(os.path.basename(filename))
        ext = os.path.splitext(filename)[1].lower()
        label = "File: " + os.path.basename(filename)
        try:
            if ext in IMAGE_EXTS:
                images.append(
                    {
                        "name": filename,
                        "mime": mime_for_image(ext),
                        "data": data,
                    }
                )
                texts.append(label + "\n[image uploaded]")
            elif ext == ".pdf":
                text = extract_pdf(data)
                texts.append(label + "\n" + (text or "[PDF had no readable text]"))
            elif ext == ".docx":
                texts.append(label + "\n" + extract_docx(data))
            elif ext in {".xlsx", ".xlsm"}:
                texts.append(label + "\n" + extract_xlsx(data))
            elif ext == ".pptx":
                texts.append(label + "\n" + extract_pptx(data))
            elif ext in TEXT_EXTS or ext == "":
                texts.append(label + "\n" + decode_bytes(data))
            else:
                texts.append(label + "\n" + decode_bytes(data))
        except Exception as error:
            raise RuntimeError("Could not read " + filename + ": " + str(error)) from error
    return texts, images, names


def openai_client():
    api_key = os.getenv("OPENAI_API_KEY", "").strip()
    if api_key.lower() in PLACEHOLDER_KEYS:
        raise RuntimeError(
            "OPENAI_API_KEY is missing. Put a real OpenAI key in the .env file."
        )
    return OpenAI(api_key=api_key)


def extract_items(syllabus, course, notes, images, default_time=""):
    client = openai_client()
    user_text = (
        "Today's date: "
        + datetime.now().strftime("%Y-%m-%d")
        + "\nCourse name: "
        + (course or "(not provided)")
        + "\nDefault due time for this course: "
        + (default_time or "(none)")
        + "\n\nStudent request: "
        + (notes or "(none — build one schedule from this input with type, date, time, and source for each item)")
        + "\n\nPasted text:\n"
        + syllabus
    )
    content = [{"type": "input_text", "text": user_text}]
    for image in images:
        encoded = base64.b64encode(image["data"]).decode("ascii")
        content.append(
            {
                "type": "input_image",
                "image_url": "data:" + image["mime"] + ";base64," + encoded,
            }
        )
    response = client.responses.create(
        model=OPENAI_MODEL,
        instructions=SYSTEM_INSTRUCTIONS,
        input=[{"role": "user", "content": content}],
        text={
            "format": {
                "type": "json_schema",
                "name": "syllabus_items",
                "schema": ITEM_SCHEMA,
                "strict": True,
            }
        },
    )
    raw_text = (response.output_text or "").strip()
    if not raw_text:
        raise RuntimeError("The AI returned an empty response.")
    try:
        parsed = json.loads(raw_text)
    except json.JSONDecodeError as error:
        raise RuntimeError("The AI did not return valid JSON.") from error
    return clean_items(parsed.get("items"), syllabus)


def request_payload():
    if request.files or request.form:
        course = str(request.form.get("course") or "").strip()
        notes = str(request.form.get("notes") or "").strip()
        origin = str(request.form.get("origin") or "Syllabus").strip() or "Syllabus"
        pasted = str(request.form.get("syllabus") or "").strip()
        default_time = str(request.form.get("defaultTime") or "").strip()
        uploaded = request.files.getlist("files") if request.files else []
        texts, images, names = read_uploaded_files(uploaded)
        chunks = []
        if pasted:
            chunks.append("Pasted text:\n" + pasted)
        chunks.extend(texts)
        syllabus = "\n\n".join(chunk for chunk in chunks if chunk).strip()
        return syllabus, course, notes, origin, images, names, default_time
    data = request.get_json(silent=True)
    if not isinstance(data, dict):
        raise ValueError("Send JSON or upload files.")
    return (
        str(data.get("syllabus") or "").strip(),
        str(data.get("course") or "").strip(),
        str(data.get("notes") or "").strip(),
        str(data.get("origin") or "Syllabus").strip() or "Syllabus",
        [],
        [],
        str(data.get("defaultTime") or "").strip(),
    )


@app.route("/")
def home():
    return send_from_directory(ROOT, "index.html")


@app.route("/baseline")
def baseline():
    return send_from_directory(ROOT, "baseline.html")


@app.route("/rules-extractor.js")
def rules_extractor():
    return send_from_directory(ROOT, "rules-extractor.js")


@app.route("/extract", methods=["POST"])
def extract():
    if rate_limited():
        return json_error("Too many requests. Please try again in a few minutes.", 429)
    try:
        syllabus, course, notes, origin, images, names, default_time = request_payload()
    except ValueError as error:
        return json_error(str(error), 400)
    except RuntimeError as error:
        return json_error(str(error), 400)
    if not syllabus and not images:
        return json_error("Paste a syllabus first.", 400)
    if len(syllabus) > MAX_TEXT_CHARS:
        syllabus = syllabus[:MAX_TEXT_CHARS]
    try:
        items = extract_items(syllabus, course, notes, images, default_time)
    except RuntimeError as error:
        return json_error(str(error), 500)          # messages written in this file, safe to show
    except Exception as error:
        app.logger.error("AI call failed: %s", type(error).__name__)   # never log pasted text or keys
        return json_error("The AI service could not finish the request. Please try again, or use the rules-only extractor.", 502)
    return jsonify(
        {
            "items": items,
            "course": course,
            "origin": origin,
            "files": names,
        }
    )


@app.errorhandler(404)
def not_found(_error):
    return json_error("That page or API route was not found.", 404)


@app.errorhandler(413)
def too_large(_error):
    return json_error("That upload is too large. Keep the total under 25 MB.", 413)


@app.errorhandler(500)
def server_error(_error):
    return json_error("Something went wrong on the server. Please try again.", 500)


if __name__ == "__main__":
    app.run(host="127.0.0.1", port=5000, debug=False)


